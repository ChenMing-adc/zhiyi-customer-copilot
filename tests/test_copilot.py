import asyncio
import json
import os
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch
from concurrent.futures import ThreadPoolExecutor

from fastapi.testclient import TestClient
from sqlalchemy import select

from app.main import create_app, ROOT
from app.copilot.importer import import_workbook, parse_workbook, ImportFailure
from app.copilot.context import get_context
from app.copilot.analysis import analyze, validate_draft
from app.copilot.store import Record, Source

SOURCE = Path(os.getenv("ZHIYI_SOURCE_PATH", ROOT.parent / "赛事方数据/赛题 1：数据共情者-业务数据.xlsx"))
EXPECTED = {"聊天记录": 998, "订单": 113, "补发换货工单": 24, "线下打款工单": 13, "物流工单": 15, "不良反应工单": 10, "售后退货工单": 18}


@unittest.skipUnless(SOURCE.is_file(), "需提供官方业务 Excel：设置 ZHIYI_SOURCE_PATH 后运行集成测试")
class CopilotTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = TemporaryDirectory()
        cls.app = create_app(Path(cls.tmp.name) / "test.sqlite3", SOURCE)
        cls.client = TestClient(cls.app)
        cls.client.__enter__()
        cls.repo = cls.app.state.store
        cls.agent = cls.login("客服 A")
        cls.agent_b = cls.login("客服 B")
        cls.manager = cls.login("主管")
        cls.admin = cls.login("管理员")

    @classmethod
    def login(cls, name):
        res = cls.client.post('/api/demo/login', json={"name": name})
        assert res.status_code == 200, res.text
        return {"Authorization": "Bearer " + res.json()["token"]}

    @classmethod
    def tearDownClass(cls):
        cls.client.__exit__(None, None, None)
        cls.tmp.cleanup()

    def context(self, sid="S00053", **params):
        res = self.client.get(f"/api/sessions/{sid}/context", params=params, headers=self.agent)
        self.assertEqual(res.status_code, 200, res.text)
        return res.json()

    def test_01_import_counts_idempotent(self):
        status = self.client.get('/api/status', headers=self.agent).json()
        self.assertIsNone(status['startup_error'])
        self.assertEqual(status['counts'], EXPECTED)
        self.assertEqual(status['sessions'], 138)
        second = import_workbook(self.repo, SOURCE)
        self.assertEqual(second['inserted'], 0)
        self.assertTrue(second['duplicate'])

    def test_02_unauthenticated_and_roles(self):
        self.assertEqual(self.client.get('/api/sessions').status_code, 401)
        self.assertEqual(self.client.post('/api/imports/official', headers=self.agent).status_code, 403)
        self.assertEqual(self.client.get('/student.html').status_code, 404)
        self.assertIn(self.client.post('/api/chat/stream', json={}).status_code, {404,405})
        self.assertEqual(self.client.post('/api/demo/login', json={"name": "主管"}, headers={"Origin": "https://untrusted.test"}).status_code, 403)

    def test_03_replay_no_future_snapshot(self):
        ctx = self.context(mode='replay', cutoff=2)
        self.assertEqual(len(ctx['messages']), 2)
        self.assertFalse(any(e['field'] in {'转账状态','工单状态','订单状态','完成时间','退款金额(元)'} for e in ctx['evidence'].values()))
        review = self.context()
        self.assertNotEqual(ctx['version'], review['version'])
        self.assertTrue(any(e['field'] == '转账状态' for e in review['evidence'].values()))
        self.assertEqual(self.client.get('/api/sessions/S00053/context?mode=replay&cutoff=999',headers=self.agent).status_code,422)

    def test_04_private_fields_and_labels_excluded(self):
        ctx = self.context()
        exposed_fields = {field for fact in ctx['facts'] for field in fact['fields']}
        exposed_fields.update(e['field'] for e in ctx['evidence'].values())
        for field in ['支付宝账号','支付宝实名','scene_major','scene_minor','is_target_buyer_message']:
            self.assertNotIn(field, exposed_fields)

    def test_05_all_sessions_and_missing_images(self):
        sessions = self.client.get('/api/sessions', headers=self.agent).json()
        no_orders = sum(not s['orders'] for s in sessions)
        self.assertEqual(no_orders,25)
        images = 0
        for s in sessions:
            ctx = get_context(self.repo, s['id'])
            for m in ctx['messages']:
                if m['image_missing']:
                    images += 1
                    self.assertEqual(m['text'],'[图片原件未提供]')
        self.assertEqual(images,29)

    def test_06_send_idempotency_and_stale_version(self):
        ctx = self.context()
        body = {"mode":"review","cutoff":None,"version":ctx['version'],"text":"我会先核实当前进度，再向您反馈。","request_id":"send-test-001"}
        response = self.client.post('/api/sessions/S00053/demo-messages',json=body,headers=self.agent)
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(self.client.post('/api/sessions/S00053/demo-messages',json=body,headers=self.agent).status_code,200)
        self.assertEqual(self.client.post('/api/sessions/S00053/demo-messages',json={**body,'request_id':'send-test-002'},headers=self.agent).status_code,409)
        self.assertEqual(self.client.post('/api/sessions/S00053/demo-messages',json={**body,'text':'另一条消息'},headers=self.agent).status_code,409)
        fresh = self.context()
        self.assertEqual(len(fresh['messages']),len(ctx['messages'])+1)
        replay = self.context(mode='replay',cutoff=ctx['total'])
        self.assertTrue(all(m['source_type']=='official_mock' for m in replay['messages']))

    def test_07_reply_commitment_blocked(self):
        ctx = self.context()
        for text in ['保证明天到账','已为您退款40元','可以为您退款','已经到账','会赔付医疗费用']:
            self.assertTrue(validate_draft(text,ctx),text)
        for text in ['请补充支付宝账号和实名信息','此前承诺不作为最终承诺']:
            self.assertTrue(validate_draft(text,ctx),text)
        self.assertFalse(validate_draft('理解您等待处理结果的焦急，我会先核实申请进度。',ctx))

    def test_08_followup_state_machine(self):
        ctx = self.context('S00268')
        body = {'session_id':'S00268','version':ctx['version'],'mode':'review','type':'不适自述跟进','content':'核实并交接专员','priority':'高','evidence_ids':[ctx['messages'][0]['id']],'due_at':'2026-05-05T10:00'}
        first = self.client.post('/api/followups',json=body,headers=self.agent)
        self.assertEqual(first.status_code,200,first.text)
        cid = first.json()['id']
        self.assertTrue(self.client.post('/api/followups',json=body,headers=self.agent).json()['duplicate'])
        def action(version, name, headers=None, **extra):
            return self.client.post(f'/api/cases/{cid}/actions',json={'version':version,'action':name,**extra},headers=headers or self.manager)
        self.assertEqual(action(1,'close',note='完成',evidence='证据').status_code,422)
        self.assertEqual(action(1,'claim',headers=self.agent).status_code,200)
        self.assertEqual(action(1,'claim',headers=self.agent_b).status_code,409)
        self.assertEqual(action(2,'note',headers=self.agent_b,note='覆盖').status_code,403)
        self.assertEqual(action(2,'close',headers=self.agent,note='完成',evidence='证据').status_code,403)
        self.assertEqual(action(2,'close',note='完成').status_code,422)
        self.assertEqual(action(2,'close',note='人工核实完成',evidence='团队自建模拟处理回执').status_code,200)
        self.assertEqual(action(3,'reopen',note='新问题',evidence='团队自建模拟回访记录').status_code,200)
        self.assertEqual(action(4,'progress').status_code,200)
        self.assertEqual(action(5,'advance').status_code,200)
        cases = self.client.get('/api/cases',headers=self.manager).json()
        case = next(c for c in cases if c['id']==cid)
        self.assertTrue(case['overdue'])
        self.assertGreaterEqual(len(case['history']),6)

    def test_09_model_structured_and_degradation(self):
        ctx = self.context('S00268')
        buyer = next(m['id'] for m in ctx['messages'] if m['role']=='买家')
        class Fake:
            calls=0
            async def stream(self,messages):
                self.calls+=1
                if self.calls==1:
                    yield json.dumps({'intent':'不良反应','emotion':'焦急','trend':'待核实','emotion_evidence':[buyer],
                        'summary':[{'text':'消费者反馈不适，待人工核实','evidence_ids':[buyer]}],'missing_info':[],'next_action':'交接专员核实'},ensure_ascii=False)
                else:
                    yield '看到您反馈的不适情况，我会整理已经提供的信息并交接专员核实，避免让您重复说明。'
        good = asyncio.run(analyze(self.repo,ctx,'test',client=Fake()))
        self.assertEqual(good['status'],'validated')
        self.assertEqual(good['calls'],2)
        class Broken:
            async def stream(self,messages):
                yield '{"emotion":"ANXIETY"}'
        bad = asyncio.run(analyze(self.repo,ctx,'test',client=Broken()))
        self.assertEqual(bad['status'],'degraded')
        self.assertTrue(bad['risks'])
        self.assertEqual(bad['draft'],'')

    def test_10_changed_source_atomic(self):
        batch, rows, counts = parse_workbook(SOURCE)
        # Exercise conflicting source protection without rewriting the source workbook.
        rows[0]['payload']['message_text'] = '不同版本内容'
        with patch('app.copilot.importer.parse_workbook',return_value=('new-batch',rows,counts)):
            with self.assertRaises(ImportFailure):
                import_workbook(self.repo,SOURCE)
        with self.repo.transaction() as db:
            self.assertIsNone(db.get(Record,'import:new-batch'))
            self.assertNotEqual(db.get(Source,rows[0]['key']).payload['message_text'],'不同版本内容')

    def test_11_bad_foreign_key_located_before_write(self):
        from openpyxl import load_workbook
        original = load_workbook(SOURCE,read_only=True,data_only=True)
        class BadSheet:
            def iter_rows(self, values_only=True):
                data = iter(original['聊天记录'].iter_rows(values_only=True))
                headers = next(data)
                yield headers
                first = list(next(data))
                first[headers.index('关联订单号')] = 'missing-order'
                yield tuple(first)
                yield from data
        class BadBook:
            sheetnames = original.sheetnames
            def __getitem__(self,key):
                return BadSheet() if key=='聊天记录' else original[key]
            def close(self):
                original.close()
        with patch('app.copilot.importer.load_workbook',return_value=BadBook()):
            with self.assertRaises(ImportFailure) as failure:
                parse_workbook(SOURCE)
        self.assertTrue(any(e['sheet']=='聊天记录' and e['row']==2 and 'missing-order' in e['error'] for e in failure.exception.errors))

    def test_12_concurrent_send_serializes_context(self):
        ctx = self.context('S00001')
        body = {'mode':'review','version':ctx['version'],'text':'我会核实您反馈的问题，再与您确认处理进度。'}
        def send(index):
            return self.client.post('/api/sessions/S00001/demo-messages',json={**body,'request_id':f'parallel-{index}'},headers=self.agent).status_code
        with ThreadPoolExecutor(max_workers=2) as pool:
            statuses = sorted(pool.map(send,[1,2]))
        self.assertEqual(statuses,[200,409])

    def test_13_wrong_emotion_evidence_degrades(self):
        ctx = self.context()
        staff = next(m['id'] for m in ctx['messages'] if m['role']=='客服')
        class Wrong:
            async def stream(self,messages):
                yield json.dumps({'intent':'退款打款','emotion':'缓和','trend':'缓和','emotion_evidence':[staff],
                    'summary':[],'missing_info':[],'next_action':'核实'})
        result = asyncio.run(analyze(self.repo,ctx,'test',client=Wrong()))
        self.assertEqual(result['status'],'degraded')
        self.assertEqual(result['draft'],'')
        self.assertEqual(result['calls'],1)

    def test_14_transport_keeps_config_and_counts_usage(self):
        import httpx
        from app.core.config import Settings
        from app.copilot.transport import ServiceAiClient
        from app.schemas.dtos import AiMessage
        settings = Settings(_env_file=None,ai_provider='openai',openai_model='qwen3.8-flash',openai_base_url='https://test.cn-beijing.maas.aliyuncs.com/compatible-mode/v1',openai_api_key='test-only')
        captured = []
        def respond(request):
            captured.append(json.loads(request.content))
            return httpx.Response(200,text='data: '+json.dumps({'choices':[{'delta':{'content':'OK'}}]})+'\n\n'+'data: '+json.dumps({'usage':{'prompt_tokens':2,'completion_tokens':1,'total_tokens':3},'choices':[]})+'\n\ndata: [DONE]\n\n')
        original_client = httpx.AsyncClient
        client = ServiceAiClient(settings)
        async def invoke():
            return ''.join([t async for t in client.stream([AiMessage(role='user',content='test')])])
        with patch('app.copilot.transport.httpx.AsyncClient',side_effect=lambda **kw:original_client(transport=httpx.MockTransport(respond))):
            self.assertEqual(asyncio.run(invoke()),'OK')
        self.assertFalse(captured[0]['enable_thinking'])
        self.assertEqual(captured[0]['max_tokens'],settings.ai_max_tokens)
        self.assertEqual(client.metrics['usage']['total_tokens'],3)

    def test_15_draft_repair_is_bounded(self):
        ctx = self.context()
        buyer = next(m['id'] for m in ctx['messages'] if m['role']=='买家')
        class Repairable:
            calls=0
            final='我会先核实补差申请的审核进度，再与您确认处理情况。'
            async def stream(self,messages):
                self.calls += 1
                if self.calls==1:
                    yield json.dumps({'intent':'退款打款','emotion':'缓和','trend':'缓和','emotion_evidence':[buyer],
                        'summary':[],'missing_info':[],'next_action':'核实进度'})
                elif self.calls==2:
                    yield '保证明天到账40元'
                else:
                    yield self.final
        fixed = asyncio.run(analyze(self.repo,ctx,'test',client=Repairable()))
        self.assertEqual(fixed['status'],'validated')
        self.assertEqual(fixed['calls'],3)
        bad = Repairable()
        bad.final = '保证明天到账40元'
        blocked = asyncio.run(analyze(self.repo,ctx,'test',client=bad))
        self.assertEqual(blocked['status'],'blocked')
        self.assertEqual(blocked['draft'],'')
        self.assertEqual(blocked['calls'],3)


if __name__=='__main__':
    unittest.main(verbosity=2)
