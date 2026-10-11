from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
import hashlib
import re

from openpyxl import load_workbook
from sqlalchemy import select

from app.copilot.store import Source, Record, audit

SHEETS = {"聊天记录": "message_id", "订单": "订单号", **{n: "工单号" for n in ["补发换货工单", "线下打款工单", "物流工单", "不良反应工单", "售后退货工单"]}}
REQUIRED = {
    "聊天记录": {"message_id", "会话ID", "消息序号", "发送时间", "角色", "店铺", "买家昵称", "message_text", "关联订单号", "关联工单号"},
    "订单": {"订单号", "会话ID", "店铺", "下单时间", "订单状态"},
}


class ImportFailure(ValueError):
    def __init__(self, errors):
        self.errors = errors
        super().__init__("导入预检失败")


def parse_workbook(path: Path):
    rows, errors, counts = [], [], {}
    with path.open("rb") as stream:
        batch = hashlib.sha256(stream.read()).hexdigest()
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet, pk in SHEETS.items():
            if sheet not in workbook.sheetnames:
                errors.append({"sheet": sheet, "row": 1, "error": "缺少工作表"})
                continue
            iterator = workbook[sheet].iter_rows(values_only=True)
            headers = [str(h or "") for h in next(iterator)]
            required = REQUIRED.get(sheet, {pk, "会话ID", "店铺", "关联订单号", "创建时间"})
            missing = required - set(headers)
            if missing:
                errors.append({"sheet": sheet, "row": 1, "error": "缺少列：" + ",".join(sorted(missing))})
                continue
            seen = set()
            counts[sheet] = 0
            for number, values in enumerate(iterator, 2):
                if not any(v is not None for v in values):
                    continue
                data = {k: (v.isoformat(sep=" ") if isinstance(v, datetime) else str(v) if v is not None else "") for k, v in zip(headers, values) if k}
                try:
                    if not data[pk] or not data["会话ID"] or not data["店铺"]:
                        raise ValueError("主键、会话ID、店铺不能为空")
                    if data[pk] in seen:
                        raise ValueError("工作表内主键重复")
                    seen.add(data[pk])
                    for field, value in data.items():
                        if value and field.endswith("时间"):
                            # Official preorder payments carry a deposit annotation.
                            # Validate the timestamp but preserve the original cell verbatim.
                            timestamp = re.sub(r"（定金）$", "", value) if field == "付款时间" else value
                            datetime.fromisoformat(timestamp)
                        if value and "(元)" in field:
                            amount = Decimal(value)
                            if not amount.is_finite():
                                raise ValueError("金额无效")
                            data[field] = format(amount.quantize(Decimal("0.01")), "f")
                    if sheet == "聊天记录":
                        int(data["消息序号"])
                        if data["角色"] not in {"买家", "客服", "系统", "系统推送"}:
                            raise ValueError("未知消息角色")
                    rows.append({"key": f"{sheet}:{data[pk]}", "sheet": sheet, "session_id": data["会话ID"], "payload": data,
                                 "provenance": {"file": path.name, "sheet": sheet, "row": number, "batch": batch, "source_type": "official_mock"}})
                    counts[sheet] += 1
                except (ValueError, InvalidOperation) as exc:
                    errors.append({"sheet": sheet, "row": number, "error": str(exc)})
        orders = {r["payload"]["订单号"]: r for r in rows if r["sheet"] == "订单"}
        tickets = {r["payload"]["工单号"]: r for r in rows if r["sheet"].endswith("工单")}
        sessions = {r["session_id"] for r in rows if r["sheet"] == "聊天记录"}
        for row in rows:
            p = row["payload"]
            issues = []
            if row["session_id"] not in sessions:
                issues.append("没有对应聊天会话")
            for field, targets in [("关联订单号", orders), ("关联工单号", tickets)]:
                ref = p.get(field)
                if not ref:
                    continue
                target = targets.get(ref)
                if not target:
                    issues.append(f"{field}引用不存在：{ref}")
                elif target["session_id"] != row["session_id"] or target["payload"]["店铺"] != p["店铺"]:
                    issues.append(f"{field}会话或店铺冲突")
            if issues:
                errors.append({"sheet": row["sheet"], "row": row["provenance"]["row"], "error": "；".join(issues)})
        if errors:
            raise ImportFailure(errors)
        return batch, rows, counts
    finally:
        workbook.close()


def import_workbook(store, path: Path):
    batch, rows, counts = parse_workbook(path)
    with store.transaction() as db:
        if db.get(Record, "import:" + batch):
            return {"counts": counts, "inserted": 0, "batch": batch, "duplicate": True}
        # Immutable source rows: a changed file must never silently replace existing evidence.
        errors = []
        for row in rows:
            existing = db.get(Source, row["key"])
            if existing and existing.payload != row["payload"]:
                errors.append({"sheet": row["sheet"], "row": row["provenance"]["row"], "error": "业务主键内容变更：本版请使用独立演示数据库导入新版本"})
        if errors:
            raise ImportFailure(errors)
        inserted = 0
        for row in rows:
            if not db.get(Source, row["key"]):
                db.add(Source(**row))
                inserted += 1
        db.add(Record(id="import:" + batch, kind="import", payload={"counts": counts, "file": path.name}))
        audit(db, "system", "import", batch, counts=counts)
    return {"counts": counts, "inserted": inserted, "batch": batch, "duplicate": False}
